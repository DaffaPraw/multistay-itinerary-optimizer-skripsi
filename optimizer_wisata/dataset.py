"""
Pemuat dataset wisata.

Menggantikan dataloader_v2 yang membaca instance Prins. Node sekarang destinasi
wisata dan akomodasi, bukan depot dan pelanggan.

Notasi mengikuti proposal:
    P, H, N      destinasi, akomodasi, gabungan keduanya
    f[i], p[h]   harga tiket masuk dan harga kamar per malam
    tau[i]       durasi kunjungan, menit
    jam[i]       jam buka dan tutup dalam menit sejak tengah malam
    s[i], u[h]   skor destinasi dan skor akomodasi, 0 sampai 10
    rho[i,d]     probabilitas presipitasi di kota i pada hari d
    s_bar[i,d]   skor destinasi sesudah dikoreksi cuaca hari d
    delta, t, g  jarak (km), waktu tempuh (menit), biaya bahan bakar antar node
"""
import math
import os
import re
from datetime import date, datetime, timedelta

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise ImportError("Paket openpyxl dibutuhkan untuk membaca dataset xlsx.") from exc


# -- pengaturan --
PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
NAMA_DATASET = "dataset_itinerary_jateng_120_lokasi_v8.xlsx"


def _cari_dataset(nama=NAMA_DATASET):
    """Urutan dicari dari yang paling dekat berkas program ke yang paling umum,
    karena folder kerja dan folder arsip menaruh dataset di tempat berbeda.
    """
    kandidat = [
        os.path.join(PROJECT_ROOT, nama),
        os.path.join(PROJECT_ROOT, os.pardir, "data", nama),
        os.path.join(PROJECT_ROOT, os.pardir, "2_dataset", nama),
        os.path.join(os.path.expanduser("~"), "Downloads", nama),
    ]
    for p in kandidat:
        if os.path.exists(p):
            return os.path.normpath(p)
    return os.path.normpath(kandidat[-1])


DATA_PATH = _cari_dataset()

SHEET_DESTINASI = "Destinations_90_Clean"
SHEET_AKOMODASI = "Accommodations_30"
SHEET_CUACA = "Weather_Features_14d"

# Persamaan (3.3). Acuan jumlah ulasan. Nilainya diambil dari jumlah ulasan
# terbesar pada dataset, sehingga lokasi terpopuler memperoleh skor popularitas
# tepat satu dan tidak ada nilai yang perlu dipotong. Konstanta di bawah hanya
# dipakai sebagai cadangan apabila acuan tidak dapat dihitung dari data.
V_REF = 70_000

# Persamaan (3.4) dan (3.5) tidak memakai bobot. Skor mutu dan skor popularitas
# digabungkan sebagai rata-rata biasa, sehingga tidak ada nilai yang perlu
# ditetapkan sendiri.

# Tabel 2.4. Pengelompokan kategori rinci dari sumber data menjadi lima kategori
# tingkat pengguna. Sumber data mencatat kategori sangat rinci sehingga sebagian
# besar kategori hanya berisi satu destinasi, dan tidak layak dipakai langsung
# sebagai batasan cakupan (2.7a) karena optimizer tidak punya alternatif.
KATEGORI_PENGGUNA = {
    "Sejarah & Museum": [
        "historical_museum", "museum_cultural_history", "royal_heritage_museum",
        "royal_palace_museum", "museum_cultural_literature", "batik_museum",
        "historical_landmark", "historical_fort", "industrial_heritage",
        "museum_special", "art_museum", "historical_heritage", "historical_park",
    ],
    "Belanja & Kuliner": [
        "shopping_mall", "shopping_market", "shopping_cultural", "traditional_market",
        "night_market", "antique_market", "culinary_market", "culinary_heritage_area",
    ],
    "Alam & Rekreasi Luar Ruang": [
        "urban_park", "nature_park", "lake_reservoir", "beach", "zoo", "waterpark",
        "theme_park_cultural",
    ],
    "Ruang Kota & Seni": [
        "urban_heritage_area", "urban_photo_spot", "city_square", "art_gallery",
        "creative_heritage", "performing_arts", "arts_cultural_center",
    ],
    "Religi": ["religious_cultural", "religious_landmark", "religious_heritage",
               "religious_pilgrimage"],
}


PETA_KATEGORI = {rinci: pengguna
                 for pengguna, daftar in KATEGORI_PENGGUNA.items()
                 for rinci in daftar}


def kategori_pengguna(kategori_sumber):
    """Kategori yang belum dipetakan dikembalikan apa adanya, bukan dibuang, supaya
    kelalaian pemetaan kelihatan di ringkasan.
    """
    return PETA_KATEGORI.get(str(kategori_sumber).strip(), str(kategori_sumber).strip())


# Persamaan (2.2). Kepekaan cuaca menurut sifat lokasi.
GAMMA = {"indoor": 0.0, "mixed": 0.5, "campuran": 0.5, "outdoor": 1.0}

# Persamaan (2.2). Saklar koreksi cuaca, 0 berarti mati dan 1 berarti nyala.
# Merupakan parameter masukan, bukan variabel keputusan, karena optimizer yang
# bebas memilih akan selalu mematikannya demi memperbaiki nilai objektif.
CUACA_AKTIF = 0

# Biaya bahan bakar, Tabel 2.9.
KBBM = 12.0       # kilometer per liter
HBBM = 10_000.0   # rupiah per liter

# Jarak masih ditaksir dari koordinat. FAKTOR_JALAN mengubah jarak garis lurus
# jadi perkiraan jarak jalan, dan kecepatannya dibedakan karena perjalanan
# dalam kota jauh lebih lambat daripada antar kota.
FAKTOR_JALAN = 1.3
KECEPATAN_KOTA_KMJAM = 25.0     # dalam satu kota
KECEPATAN_ANTARKOTA_KMJAM = 55.0

RADIUS_BUMI_KM = 6371.0


def _teks(nilai):
    return "" if nilai is None else str(nilai).strip()


def _angka(nilai, nama_field, node_id, wajib=True, bawaan=0.0):
    teks = _teks(nilai)
    if teks == "":
        if wajib:
            raise ValueError(f"{nama_field} kosong pada node {node_id}.")
        return bawaan
    try:
        return float(teks)
    except ValueError as exc:
        raise ValueError(f"{nama_field} bukan angka pada node {node_id}: {nilai!r}") from exc


def _haversine_km(a, b):
    lat1, lon1 = math.radians(a[0]), math.radians(a[1])
    lat2, lon2 = math.radians(b[0]), math.radians(b[1])
    dlat, dlon = lat2 - lat1, lon2 - lon1
    h = math.sin(dlat / 2) ** 2 + math.cos(lat1) * math.cos(lat2) * math.sin(dlon / 2) ** 2
    return 2 * RADIUS_BUMI_KM * math.asin(math.sqrt(h))


def _baca_sheet(wb, nama):
    if nama not in wb.sheetnames:
        raise ValueError(f"Sheet '{nama}' tidak ada pada dataset.")
    ws = wb[nama]
    baris = ws.iter_rows(values_only=True)
    try:
        header = [_teks(sel) for sel in next(baris)]
    except StopIteration:
        return []
    hasil = []
    for nilai in baris:
        if nilai is None or all(sel is None for sel in nilai):
            continue
        hasil.append({header[i]: nilai[i] for i in range(min(len(header), len(nilai)))})
    return hasil


# -- Persamaan (3.1)-(3.5) --
def rating_terskala(rating, skala):
    """Persamaan (3.1). Rating ke skala lima.

    Sumber memakai skala 5 dan 10 bercampur, jadi perlu diseragamkan dulu.
    Identitas untuk sumber yang sudah berskala lima.
    """
    if skala <= 0:
        raise ValueError("Skala rating harus positif.")
    return 5.0 * rating / skala


def skor_mutu(rating_skala_lima):
    """Persamaan (3.2). Rating skala lima menjadi skor mutu 0 sampai 1."""
    return min(max((rating_skala_lima - 1.0) / 4.0, 0.0), 1.0)


def skor_popularitas(jumlah_ulasan, v_ref=None):
    """Persamaan (3.3). Jumlah ulasan jadi skor popularitas 0..1.

    Logaritma karena sebaran jumlah ulasan sangat lebar. Dipotong di 1 supaya
    lokasi yang melampaui V_ref tidak keluar rentang.
    """
    if jumlah_ulasan < 0:
        raise ValueError("Jumlah ulasan tidak boleh negatif.")
    if v_ref is None:
        v_ref = V_REF
    return min(math.log(jumlah_ulasan + 1.0) / math.log(v_ref + 1.0), 1.0)


def skor_gabungan(mu, pi):
    """Persamaan (3.4) dan (3.5). Rata-rata mutu dan popularitas, 0..10.

    Sengaja tanpa bobot supaya tidak ada nilai yang perlu dipertanggungjawabkan.
    """
    return 5.0 * (mu + pi)


def skor_terkoreksi(skor, gamma, rho, w=CUACA_AKTIF):
    """Persamaan (2.2). Skor dikoreksi probabilitas presipitasi.

    w adalah saklar; w=0 berarti skor kembali ke nilai dasarnya.
    """
    return skor * (1.0 - w * gamma * rho)


_POLA_JAM = re.compile(r"(\d{1,2}):(\d{2})\s*-\s*(\d{1,2}):(\d{2})")


def parse_jam_operasional(teks, bawaan=(8 * 60, 21 * 60)):
    """Jam buka dan tutup dalam menit sejak tengah malam.

    Format gaya OpenStreetMap, misalnya 'Mo-Su 07:00-21:00'. Tanpa jam eksplisit
    lokasi dianggap buka sepanjang hari aktivitas; nilai balik ketiga menandai
    bahwa itu asumsi, bukan data.
    """
    cocok = _POLA_JAM.search(teks or "")
    if not cocok:
        return bawaan[0], bawaan[1], True
    buka = int(cocok.group(1)) * 60 + int(cocok.group(2))
    tutup = int(cocok.group(3)) * 60 + int(cocok.group(4))
    if tutup <= buka:      # melewati tengah malam, dipotong di akhir hari
        tutup = 24 * 60
    return buka, tutup, False


def _acuan_ulasan(baris):
    """Jumlah ulasan terbesar di seluruh dataset, dipakai sebagai V acuan (3.3)."""
    nilai = []
    for b in baris:
        if _teks(b.get("usable_for_model")).lower() in ("0", "false", "no"):
            continue
        teks = _teks(b.get("review_count"))
        if teks:
            nilai.append(float(teks))
    return max(nilai) if nilai else float(V_REF)


def _muat_lokasi(baris, peran, v_ref=None):
    lokasi = {}
    for b in baris:
        node_id = _teks(b.get("location_id"))
        if not node_id:
            continue
        if _teks(b.get("usable_for_model")).lower() in ("0", "false", "no"):
            continue

        rating = _angka(b.get("rating_value"), "rating_value", node_id)
        skala = _angka(b.get("rating_scale"), "rating_scale", node_id)
        ulasan = _angka(b.get("review_count"), "review_count", node_id)

        mu = skor_mutu(rating_terskala(rating, skala))
        pi = skor_popularitas(ulasan, v_ref)

        sifat = _teks(b.get("indoor_outdoor")).lower()
        lokasi[node_id] = {
            "id": node_id,
            "nama": _teks(b.get("name")),
            "peran": peran,
            "kota": _teks(b.get("hub_city")),
            "kategori_sumber": _teks(b.get("category")),
            "kategori": kategori_pengguna(_teks(b.get("category"))),
            "tier": _teks(b.get("accommodation_tier")),
            "lat": _angka(b.get("latitude"), "latitude", node_id),
            "lon": _angka(b.get("longitude"), "longitude", node_id),
            "jam_buka_teks": _teks(b.get("opening_hours")),
            "tiket": _angka(b.get("ticket_price_idr"), "ticket_price_idr", node_id, wajib=False),
            "harga_malam": _angka(b.get("price_per_night_idr"), "price_per_night_idr", node_id, wajib=False),
            "durasi": _angka(b.get("visit_duration_min"), "visit_duration_min", node_id, wajib=False),
            "sifat": sifat,
            "gamma": GAMMA.get(sifat, 0.5),
            "rating": rating,
            "skala": skala,
            "ulasan": int(ulasan),
            "mu": mu,
            "pi": pi,
            "skor": skor_gabungan(mu, pi),
            "sumber_resmi": _teks(b.get("official_source_url")),
            "tier_resmi": _teks(b.get("official_source_tier")),
        }
    return lokasi


def _muat_cuaca(baris):
    """Cuaca per jam jadi satu probabilitas presipitasi per kota per tanggal.

    Rata-rata diambil pada jam aktivitas saja, bukan 24 jam, karena wisatawan
    tidak di luar sepanjang malam.
    """
    kumpul = {}
    for b in baris:
        kota = _teks(b.get("hub_city"))
        nilai_tanggal = b.get("date")
        if isinstance(nilai_tanggal, datetime):
            tanggal = nilai_tanggal.date()
        elif isinstance(nilai_tanggal, date):
            tanggal = nilai_tanggal
        else:
            teks = _teks(nilai_tanggal)[:10]
            if not teks:
                continue
            try:
                tanggal = datetime.strptime(teks, "%Y-%m-%d").date()
            except ValueError:
                continue
        jam = b.get("hour")
        jam = int(jam) if jam is not None and _teks(jam) != "" else 12
        if not 6 <= jam <= 21:
            continue
        prob = _angka(b.get("probability_precipitation"), "probability_precipitation",
                      kota, wajib=False)
        kumpul.setdefault((kota, tanggal), []).append(prob / 100.0)
    return {kunci: sum(nilai) / len(nilai) for kunci, nilai in kumpul.items() if nilai}


def _kecepatan(kota_u, kota_v):
    return KECEPATAN_KOTA_KMJAM if kota_u == kota_v else KECEPATAN_ANTARKOTA_KMJAM


def bangun_matriks(node, lapor=False):
    """Menyusun delta, t, dan g untuk seluruh pasangan node.

    Jarak ditaksir dari koordinat memakai haversine dikali FAKTOR_JALAN, bukan
    diambil dari sheet Route_Matrix pada dataset, karena matriks di sheet itu
    sendiri masih taksiran dan belum memuat seluruh node. Nilainya perlu diganti
    hasil layanan rute sebelum percobaan akhir.

    Taksiran ini simetris, jadi hanya pasangan dengan u lebih kecil daripada v
    yang dihitung lalu dicerminkan. Untuk 120 node perhitungan haversine turun
    dari 14.280 jadi 7.140. Waktu tempuh ikut simetris karena hanya bergantung
    jarak dan kelas kecepatan antar kota.

    Pasangannya terbagi tiga blok: antar destinasi, antar akomodasi, dan
    destinasi dengan akomodasi. Ketiganya dipakai model, karena rute harian
    berangkat dari akomodasi, melewati destinasi, lalu berakhir di akomodasi,
    dan hari tanpa kunjungan menghasilkan arc dari akomodasi ke akomodasi.
    """
    ids = list(node.keys())
    delta, t, g = {}, {}, {}
    hitung = 0

    for u in ids:
        delta[u, u] = t[u, u] = g[u, u] = 0.0

    for a in range(len(ids)):
        for b in range(a + 1, len(ids)):
            u, v = ids[a], ids[b]
            km = _haversine_km((node[u]["lat"], node[u]["lon"]),
                               (node[v]["lat"], node[v]["lon"])) * FAKTOR_JALAN
            menit = 60.0 * km / _kecepatan(node[u]["kota"], node[v]["kota"])
            rupiah = km / KBBM * HBBM
            hitung += 1
            # dicerminkan, bukan dihitung ulang
            delta[u, v] = delta[v, u] = km
            t[u, v] = t[v, u] = menit
            g[u, v] = g[v, u] = rupiah

    if lapor:
        n = len(ids)
        print(f"  matriks rute: {n} node, {hitung:,} perhitungan haversine "
              f"untuk {n * (n - 1):,} pasangan berarah "
              f"({1 - hitung / (n * (n - 1)):.0%} dihemat lewat simetri)")
    return delta, t, g


def validasi_dataset(P, H, node, cuaca, tanggal):
    """Mengumpulkan seluruh masalah lebih dulu, baru melempar sekali."""
    galat = []
    if not P:
        galat.append("Tidak ada destinasi kandidat setelah penyaringan.")
    if not H:
        galat.append("Tidak ada akomodasi kandidat setelah penyaringan.")

    for i in P:
        if node[i]["durasi"] <= 0:
            galat.append(f"Durasi kunjungan kosong atau nol pada destinasi {i}.")
        if node[i]["skor"] < 0 or node[i]["skor"] > 10:
            galat.append(f"Skor destinasi {i} di luar rentang 0 sampai 10: {node[i]['skor']:.4f}.")
    for h in H:
        if node[h]["harga_malam"] <= 0:
            galat.append(f"Harga akomodasi {h} kosong atau nol.")
        if node[h]["skor"] < 0 or node[h]["skor"] > 10:
            galat.append(f"Skor akomodasi {h} di luar rentang 0 sampai 10: {node[h]['skor']:.4f}.")

    for kota in {node[i]["kota"] for i in P}:
        if not any(node[h]["kota"] == kota for h in H):
            galat.append(f"Kota {kota} punya destinasi tetapi tidak punya akomodasi kandidat.")

    # rho dibangun dengan indeks langsung sesudah validasi ini. Alternatifnya
    # memberi nilai bawaan 0.0, tapi 0.0 berarti cerah, jadi data yang bolong
    # akan tampak seperti cuaca bagus.
    hilang = sorted({(node[i]["kota"], t) for i in P for t in tanggal
                     if (node[i]["kota"], t) not in cuaca})
    if hilang:
        galat.append(
            f"Data cuaca tidak ada untuk {len(hilang)} pasangan kota dan tanggal, contoh {hilang[:3]}."
        )

    if galat:
        raise ValueError("Validasi dataset gagal:\n- " + "\n- ".join(galat))


def muat_dataset(data_path=DATA_PATH, kota_dipilih=None, kategori_dipilih=None,
                 tanggal_mulai=None, jumlah_hari=3, v_ref=None,
                 cuaca_aktif=CUACA_AKTIF):
    """Himpunan node beserta parameter model.

    Tanggal di luar rentang data cuaca digeser ke awal rentang, dan
    pergeserannya dicatat di 'catatan' supaya tidak diam-diam.
    """
    if not os.path.exists(data_path):
        raise FileNotFoundError(f"Dataset tidak ditemukan: {data_path}")

    wb = load_workbook(data_path, data_only=True, read_only=True)
    try:
        baris_destinasi = _baca_sheet(wb, SHEET_DESTINASI)
        baris_akomodasi = _baca_sheet(wb, SHEET_AKOMODASI)
        baris_cuaca = _baca_sheet(wb, SHEET_CUACA)
    finally:
        wb.close()

    if v_ref is None:
        v_ref = _acuan_ulasan(baris_destinasi + baris_akomodasi)

    destinasi = _muat_lokasi(baris_destinasi, "destinasi", v_ref)
    akomodasi = _muat_lokasi(baris_akomodasi, "akomodasi", v_ref)
    cuaca = _muat_cuaca(baris_cuaca)
    catatan = []

    if kota_dipilih:
        kota_dipilih = set(kota_dipilih)
        destinasi = {k: v for k, v in destinasi.items() if v["kota"] in kota_dipilih}
        akomodasi = {k: v for k, v in akomodasi.items() if v["kota"] in kota_dipilih}
    if kategori_dipilih:
        kategori_dipilih = set(kategori_dipilih)
        destinasi = {k: v for k, v in destinasi.items() if v["kategori"] in kategori_dipilih}

    node = {**destinasi, **akomodasi}
    P = sorted(destinasi)
    H = sorted(akomodasi)

    tanggal_tersedia = sorted({t for (_, t) in cuaca})
    if not tanggal_tersedia:
        raise ValueError("Sheet cuaca tidak memuat satu pun tanggal yang terbaca.")
    if tanggal_mulai is None:
        tanggal_mulai = tanggal_tersedia[0]
    elif isinstance(tanggal_mulai, str):
        tanggal_mulai = datetime.strptime(tanggal_mulai, "%Y-%m-%d").date()

    tanggal = [tanggal_mulai + timedelta(days=k) for k in range(jumlah_hari)]
    if any(t not in tanggal_tersedia for t in tanggal):
        catatan.append(
            f"Tanggal {tanggal[0]} sampai {tanggal[-1]} di luar rentang data cuaca "
            f"{tanggal_tersedia[0]} sampai {tanggal_tersedia[-1]}. Digeser ke awal rentang."
        )
        tanggal_mulai = tanggal_tersedia[0]
        tanggal = [tanggal_mulai + timedelta(days=k) for k in range(jumlah_hari)]

    validasi_dataset(P, H, node, cuaca, tanggal)

    rho = {}
    for i in P:
        for d, tgl in enumerate(tanggal, start=1):
            rho[i, d] = cuaca[node[i]["kota"], tgl]

    jam = {}
    asumsi_jam = []
    for i in P:
        buka, tutup, diasumsikan = parse_jam_operasional(node[i]["jam_buka_teks"])
        jam[i] = (buka, tutup)
        if diasumsikan:
            asumsi_jam.append(i)
    if asumsi_jam:
        catatan.append(
            f"{len(asumsi_jam)} destinasi tidak punya jam operasional eksplisit dan "
            f"dianggap buka sepanjang hari aktivitas."
        )

    delta, t_mat, g = bangun_matriks(node)

    s_bar = {
        (i, d): skor_terkoreksi(node[i]["skor"], node[i]["gamma"], rho[i, d], cuaca_aktif)
        for i in P
        for d in range(1, jumlah_hari + 1)
    }

    return {
        "data_path": data_path,
        "P": P,
        "H": H,
        "N": P + H,
        "node": node,
        "m": jumlah_hari,
        "D": list(range(1, jumlah_hari + 1)),
        "tanggal": tanggal,
        "kota": sorted({node[i]["kota"] for i in P}),
        "kategori": sorted({node[i]["kategori"] for i in P}),
        "s": {i: node[i]["skor"] for i in P},
        "u": {h: node[h]["skor"] for h in H},
        "s_bar": s_bar,
        "rho": rho,
        "gamma": {i: node[i]["gamma"] for i in P},
        "f": {i: node[i]["tiket"] for i in P},
        "p": {h: node[h]["harga_malam"] for h in H},
        "tau": {i: node[i]["durasi"] for i in P},
        "jam": jam,
        "delta": delta,
        "t": t_mat,
        "g": g,
        "v_ref": v_ref,
        "cuaca_aktif": cuaca_aktif,
        "catatan": catatan,
    }


def cetak_ringkasan(data_path=DATA_PATH):
    data = muat_dataset(data_path)
    node = data["node"]
    print("Dataset itinerary wisata multi-menginap")
    print("=" * 78)
    print(f"berkas            : {data['data_path']}")
    print(f"destinasi         : {len(data['P'])}")
    print(f"akomodasi         : {len(data['H'])}")
    print(f"kota              : {', '.join(data['kota'])}")
    print(f"kategori          : {len(data['kategori'])}")
    print(f"hari perjalanan   : {data['m']}  ({data['tanggal'][0]} sampai {data['tanggal'][-1]})")
    print(f"V_ref             : {data['v_ref']:,.0f}")
    print("penggabungan skor : rata-rata mutu dan popularitas, tanpa bobot")

    s, u = data["s"], data["u"]
    print(f"skor destinasi    : min {min(s.values()):.2f}  maks {max(s.values()):.2f}  "
          f"rata {sum(s.values()) / len(s):.2f}")
    print(f"skor akomodasi    : min {min(u.values()):.2f}  maks {max(u.values()):.2f}  "
          f"rata {sum(u.values()) / len(u):.2f}")

    rho = data["rho"]
    print(f"saklar cuaca w    : {data['cuaca_aktif']}")
    print(f"presipitasi rata  : {sum(rho.values()) / len(rho):.3f}  maks {max(rho.values()):.3f}")
    s_bar = data["s_bar"]
    print(f"skor terkoreksi   : rata {sum(s_bar.values()) / len(s_bar):.2f}  (w={data['cuaca_aktif']})")
    print(f"faktor jalan      : {FAKTOR_JALAN} kali jarak haversine")
    delta = data["delta"]
    dalam = [v for (u, w), v in delta.items()
             if u != w and node[u]["kota"] == node[w]["kota"]]
    antar = [v for (u, w), v in delta.items()
             if u != w and node[u]["kota"] != node[w]["kota"]]
    print(f"jarak dalam kota  : rata {sum(dalam) / len(dalam):5.1f} km  "
          f"maks {max(dalam):5.1f} km  ({len(dalam):,} pasangan)")
    print(f"jarak antar kota  : rata {sum(antar) / len(antar):5.1f} km  "
          f"maks {max(antar):5.1f} km  ({len(antar):,} pasangan)")
    g = data["g"]
    print(f"biaya bahan bakar : rata Rp{sum(g.values()) / len(g):,.0f} per arc  "
          f"maks Rp{max(g.values()):,.0f}")

    f, p, tau = data["f"], data["p"], data["tau"]
    print(f"tiket masuk       : Rp{min(f.values()):,.0f} sampai Rp{max(f.values()):,.0f}")
    print(f"harga kamar       : Rp{min(p.values()):,.0f} sampai Rp{max(p.values()):,.0f}")
    print(f"durasi kunjungan  : {min(tau.values()):.0f} sampai {max(tau.values()):.0f} menit")

    print("\njumlah destinasi per kategori")
    for k in data["kategori"]:
        n = sum(1 for i in data["P"] if node[i]["kategori"] == k)
        print(f"  {k:30} {n:3}")

    print("\nlima destinasi berskor tertinggi")
    for i in sorted(data["P"], key=lambda k: -s[k])[:5]:
        print(f"  {i}  {node[i]['nama'][:34]:34}  skor {s[i]:5.2f}  "
              f"mu {node[i]['mu']:.3f}  pi {node[i]['pi']:.3f}  ulasan {node[i]['ulasan']:,}")

    print("\nlima destinasi berskor terendah")
    for i in sorted(data["P"], key=lambda k: s[k])[:5]:
        print(f"  {i}  {node[i]['nama'][:34]:34}  skor {s[i]:5.2f}  "
              f"mu {node[i]['mu']:.3f}  pi {node[i]['pi']:.3f}  ulasan {node[i]['ulasan']:,}")

    if data["catatan"]:
        print("\ncatatan")
        for c in data["catatan"]:
            print(f"  - {c}")


if __name__ == "__main__":
    cetak_ringkasan()
